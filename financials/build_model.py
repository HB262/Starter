"""Builds Jet_Business_Financial_Model.xlsx.

Run:  python3 financials/build_model.py
Every number in the workbook is either a labeled input on the Assumptions /
Fleet Plan sheets (blue text) or a formula that flows from them.

Three growth strategies (Ownership, Management-led, Hybrid) each get their own
Forecast and Debt sheet built from the same formulas, so they can be compared
side by side on the Dashboard. The Downside / Base / Upside scenario switch on
Assumptions!C3 drives the economics for all three at once.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = "financials/Jet_Business_Financial_Model.xlsx"
YEARS = list(range(2027, 2037))
FIRST_COL = 3  # column C holds the first forecast year
YCOLS = [get_column_letter(FIRST_COL + i) for i in range(len(YEARS))]
LASTY = YCOLS[-1]
MAX_SLOTS = 7  # owned-aircraft slots in each debt schedule

FONT = "Arial"
BLUE = Font(name=FONT, size=10, color="0000FF")
BLACK = Font(name=FONT, size=10, color="000000")
GREEN = Font(name=FONT, size=10, color="008000")
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
SUB = Font(name=FONT, size=10, italic=True, color="595959")
HDR_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
TOP = Border(top=Side(style="thin", color="000000"))

USD = '$#,##0;($#,##0);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.0"x"'
YEAR = "0"

STRATEGIES = [
    # code, display name, sheet suffix, description
    ("own", "Ownership", "Own", "Buy more jets; own certificate operates only owned jets"),
    ("mgmt", "Management-led", "Mgmt", "Keep one owned jet; grow by managing other owners' jets"),
    ("hyb", "Hybrid", "Hybrid", "Add owned jets slowly and grow third-party management"),
]
FLEET = {
    "own": {"owned": [1, 1, 2, 2, 3, 3, 4, 4, 5, 5], "tp": [0] * 10, "cert": 2030},
    "mgmt": {"owned": [1] * 10, "tp": [0, 0, 0, 2, 4, 6, 9, 12, 15, 18], "cert": 2030},
    "hyb": {"owned": [1, 1, 1, 2, 2, 2, 3, 3, 3, 3], "tp": [0, 0, 0, 1, 3, 5, 7, 9, 11, 12], "cert": 2030},
}

wb = Workbook()


def setup(ws, title, subtitle, widths):
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = subtitle
    ws["A2"].font = SUB
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False


def header(ws, row, values, start_col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start_col + i, value=v)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center" if i else "left", wrap_text=True)


def section(ws, row, text, ncols):
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = SEC_FILL
    ws.cell(row=row, column=1, value=text).font = BOLD


def put(ws, ref, value, fmt=None, font=None, fill=None, bold=False):
    c = ws[ref]
    c.value = value
    if font is None:
        if isinstance(value, str) and value.startswith("="):
            font = GREEN if "!" in value else BLACK
        else:
            font = BLUE
    if bold:
        font = Font(name=font.name, size=font.size, color=font.color, bold=True)
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def q(sheet):
    return f"'{sheet}'"


# ---------------------------------------------------------------- Assumptions
A = wb.active
A.title = "Assumptions"
setup(A, "Assumptions: Citation Sovereign+ NetJets overflow business (EJM certificate)",
      "Blue = input you can change. Yellow = key lever. Column F (Live) feeds every forecast "
      "based on the scenario selected in C3.",
      {"A": 52, "B": 12, "C": 14, "D": 14, "E": 14, "F": 14, "G": 90})
A["A3"] = "Active scenario (1 = Downside, 2 = Base, 3 = Upside)"
A["A3"].font = BOLD
put(A, "C3", 2, fill=KEY_FILL)
put(A, "D3", '=INDEX({"Downside","Base","Upside"},$C$3)')
header(A, 5, ["Assumption", "Unit", "Downside", "Base", "Upside", "Live", "Source / note"])

# (key, label, unit, down, base, up, fmt, note, key_lever)
ASSUMPTIONS = [
    ("SEC", "Aircraft acquisition (per aircraft)"),
    ("price", "Purchase price, lightly used Citation Sovereign+ (2014-2016)", "$", 12_500_000, 11_500_000, 10_500_000, USD,
     "Chief Pilot (Derek) prefers a lightly used Sovereign+, e.g. a 2015 on Controller.com. Guardian Jet shows 2015-2018 "
     "Sovereign+ asking $10.3-15.0M (avg $12.6M); a 2015 with ~2,000 hrs is listed 'make offer'. Asking > sale price. "
     "VERIFY with a broker/JETNET before LOI.", True),
    ("price_chg", "Annual change in price paid for later aircraft", "%", 0.0, 0.0, 0.0, PCT,
     "Assumes later purchases are similar-vintage aircraft at today's price. Adjust if the market moves.", False),
    ("prebuy", "Pre-buy inspection (Textron service center)", "$", 85_000, 75_000, 60_000, USD,
     "Estimate; includes ferry to inspection facility and borescope.", False),
    ("legal", "Closing legal, escrow, title, registration", "$", 60_000, 50_000, 40_000, USD, "Estimate.", False),
    ("onboard", "Part 135 onboarding, conformity, Wi-Fi/avionics updates", "$", 200_000, 150_000, 100_000, USD,
     "Adding aircraft to operator's certificate (conformity, manuals, OpSpecs), plus connectivity charter "
     "customers expect. Estimate.", False),
    ("refurb", "Cosmetic interior / paint refresh", "$", 350_000, 250_000, 150_000, USD,
     "Fractional customers expect fleet-standard cabins. Estimate.", False),
    ("init_train", "Initial crew training and hiring (per aircraft)", "$", 100_000, 75_000, 50_000, USD,
     "Type rating for new hires (CE-680) and Part 135 initial at EJM's FlightSafety rates. Estimate.", False),
    ("squawk", "Post-purchase squawk and unscheduled maintenance reserve (one-time)", "$", 500_000, 350_000, 250_000, USD,
     "Deliberately conservative (Derek): covers pre-buy findings not credited by the seller, deferred items, and "
     "first-year surprises on a 10-year-old airframe. Funded at closing; unspent cash stays in the company.", False),
    ("entity", "Company formation, legal, accounting, insurance setup (once)", "$", 75_000, 50_000, 35_000, USD,
     "LLC formation, operating agreement, aviation counsel, CPA, insurance broker. Estimate.", False),
    ("offering", "Investor offering costs (securities counsel, PPM, Form D) (once)", "$", 100_000, 75_000, 50_000, USD,
     "Reg D private placement to accredited investors. Estimate.", False),
    ("SEC", "Financing"),
    ("ltv", "Loan-to-value (debt share of purchase price)", "%", 0.80, 0.80, 0.80, PCT,
     "Henry: assume at least 80% debt. Without a strong personal guarantor, lenders may cap an investor-owned SPV at "
     "60-70%. See 'Financing Scenarios'.", True),
    ("rate", "Interest rate (fixed)", "%", 0.0800, 0.0725, 0.0675, PCT,
     "Advertised jet loan rates ~5.6-9% (Sep/Oct 2026 rate sheets); prime 7.00% (Oct 2026). flycraft.com/aircraft-financing/aircraft-loan-rates", False),
    ("amort", "Amortization period", "years", 15, 15, 20, NUM,
     "Typical 10-20 yr amortization, often with a 5-7 yr balloon refinanced at the time. Model assumes refinance on same terms.", False),
    ("SEC", "Charter revenue (per owned aircraft)"),
    ("hours", "Billable revenue hours per aircraft at maturity", "hrs/yr", 500, 600, 700, NUM,
     "Includes billable repositioning. Dedicated fractional supplemental lift can run 800+ hrs; retail-only charter often 300-450.", True),
    ("months_y1", "Months in revenue service in 2027", "months", 5, 6, 6, NUM,
     "Henry: target revenue flying mid-2027.", False),
    ("ramp1", "Utilization in first year of service (% of mature)", "%", 0.50, 0.60, 0.70, PCT,
     "Ramp while operator builds the aircraft into its overflow rotation.", False),
    ("ramp2", "Utilization in second year (% of mature)", "%", 0.80, 0.90, 1.00, PCT, "", False),
    ("mix", "Share of hours flown for NetJets (via EJM)", "%", 0.95, 0.95, 0.95, PCT,
     "Derek: EJM's managed fleet flies ~95% NetJets customers, so little charter selling is needed. The rest is retail charter. "
     "Model the NetJets cut-off risk with the hours sensitivity on Unit Economics.", True),
    ("rate_ovf", "Revenue per billable hour, NetJets trips (via EJM)", "$/hr", 5_750, 6_000, 6_250, USD,
     "Henry (Oct 31, 2026): plan on $6,000/hr. Treated as all-in (wet): fuel is paid out of it unless the fuel line below says otherwise. Confirm with EJM.", True),
    ("rate_ret", "Revenue per billable hour, retail charter", "$/hr", 5_750, 6_000, 6_250, USD,
     "Henry: plan on $6,000/hr for all revenue. Excludes 7.5% federal excise tax.", True),
    ("rate_g", "Annual charter rate growth", "%", 0.020, 0.030, 0.035, PCT, "", False),
    ("ferry", "Unpaid positioning hours (% of billable)", "%", 0.15, 0.10, 0.08, PCT,
     "Denver base means some empty legs to the coasts. Estimate.", False),
    ("SEC", "Variable operating costs (per flight hour)"),
    ("gph", "Fuel burn", "gal/hr", 290, 274, 265, NUM, "Guardian Jet Sovereign+ operating cost sheet (300 hrs/yr), supplied by Henry Oct 31, 2026: 274 gal/hr.", False),
    ("fuel", "Jet-A price", "$/gal", 7.75, 7.54, 7.00, '$#,##0.00',
     "Guardian Jet Sovereign+ operating cost sheet (300 hrs/yr), supplied by Henry Oct 31, 2026: $7.54/gal. EJM's negotiated price may be lower; not credited in base.", False),
    ("fuel_paid", "Share of fuel cost paid separately by NetJets / the customer", "%", 0.0, 0.0, 0.0, PCT,
     "0% = the $6,000/hr rate is wet and we buy the fuel. Set to 100% if EJM confirms NetJets reimburses fuel.", True),
    ("maint", "Airframe maintenance (5-year average)", "$/hr", 1_340, 1_220, 1_160, USD,
     "Guardian Jet Sovereign+ operating cost sheet (300 hrs/yr), supplied by Henry Oct 31, 2026: $1,220/hr. Downside +10%. Maintenance contingency below sits on top (Derek: keep it conservative).", False),
    ("engine", "Engine program (2 x PW306D)", "$/hr", 1_100, 1_017, 906, USD,
     "Guardian Jet Sovereign+ operating cost sheet (300 hrs/yr), supplied by Henry Oct 31, 2026: $1,017/hr. A 2015 listing shows PowerAdvantage+ at $453/engine/hr ($906/hr).", False),
    ("apu", "APU program (not on the Guardian sheet)", "$/hr", 95, 0, 0, USD, "Assumed inside Guardian's maintenance figure and the maintenance contingency; downside adds a separate program.", False),
    ("maint_cont", "Maintenance contingency added to airframe, engine and APU costs", "%", 0.20, 0.15, 0.10, PCT,
     "Extra cushion on top of the benchmarks for an out-of-production airframe: AOG parts, corrosion, gear and "
     "landing-gear overhauls, inspection overruns. Derek asked for very conservative maintenance budgets.", False),
    ("trip", "Crew travel, landing, handling fees (not on the Guardian sheet)", "$/hr", 500, 450, 400, USD,
     "Part 135 trip costs not rebilled. Flycraft range $400-650/hr. Confirm with EJM what NetJets pays for.", False),
    ("SEC", "Fixed operating costs (per owned aircraft, per year)"),
    ("captains", "Captains per aircraft", "#", 2, 2, 2, NUM, "Chief pilot (Henry's partner) is one of the captains.", False),
    ("fos", "Second PIC-rated pilots per aircraft", "#", 1, 1, 1, NUM,
     "Derek: both seats hold full PIC type ratings (NetJets standard). Plus contract pilots for peaks.", False),
    ("cap_sal", "Captain salary", "$", 285_000, 265_000, 250_000, USD,
     "Set so crew pay, benefits and contract pilots total about $1.13M, matching the Guardian Jet Sovereign+ operating cost sheet (300 hrs/yr), supplied by Henry Oct 31, 2026.", False),
    ("fo_sal", "Second pilot salary (PIC type-rated)", "$", 255_000, 240_000, 225_000, USD,
     "Derek: both seats hold full PIC type ratings. Calibrated to the Guardian crew total.", False),
    ("benefits", "Benefits and payroll tax load", "%", 0.32, 0.30, 0.28, PCT, "Calibrated to the Guardian crew total.", False),
    ("contract", "Contract pilot budget", "$", 140_000, 125_000, 110_000, USD, "Calibrated to the Guardian crew total.", False),
    ("train", "Recurrent training per pilot per year (FlightSafety every 6 months)", "$", 55_000, 45_000, 40_000, USD,
     "Derek: pilots return to FlightSafety every 6 months at EJM's training rates. Two sessions plus travel per pilot.", False),
    ("hangar", "Hangar", "$", 90_000, 78_230, 66_000, USD,
     "Guardian Jet Sovereign+ operating cost sheet (300 hrs/yr), supplied by Henry Oct 31, 2026: $78,230. EJM has hangars; confirm KBJC space.", False),
    ("ins", "Insurance (hull + liability)", "$", 85_000, 53_968, 50_000, USD,
     "Guardian Jet Sovereign+ operating cost sheet (300 hrs/yr), supplied by Henry Oct 31, 2026: hull $18,968 + liability $35,000. Charter use and NetJets limits may cost more: get a quote.", False),
    ("mgmt_fee", "EJM management fee (while on EJM's certificate)", "$/month", 11_000, 10_000, 8_500, USD,
     "Covers EJM's Part 135 certificate, OpSpecs, admin and NetJets compliance. Light jet ~$3K/month to heavy $10K+ "
     "(Jet Linx). Get EJM's actual fee.", True),
    ("comm", "EJM commission on charter revenue", "%", 0.12, 0.10, 0.08, PCT,
     "Commonly 10-15% of charter revenue. Confirm whether EJM takes a cut of NetJets trips.", True),
    ("other", "Other fixed (Guardian 'miscellaneous' less training)", "$", 400_000, 363_997, 340_000, USD,
     "Guardian Jet Sovereign+ operating cost sheet (300 hrs/yr), supplied by Henry Oct 31, 2026: miscellaneous $498,997, less the $135K of recurrent training shown on its own line.", False),
    ("conting", "Operating contingency (AOG downtime, inspection overruns, surprises)", "$", 200_000, 150_000, 100_000, USD,
     "Deliberately conservative (Derek). Sits on top of the maintenance contingency.", False),
    ("ga", "Company G&A before own certificate (incl. investor reporting)", "$", 120_000, 100_000, 85_000, USD,
     "Accounting, K-1s and investor reporting, legal, owner-level insurance, chief pilot management stipend.", False),
    ("infl", "Annual cost inflation", "%", 0.040, 0.030, 0.025, PCT, "", False),
    ("wc_months", "Cash reserve to hold (months of fixed cost per aircraft)", "months", 6, 6, 5, NUM,
     "Lenders and management companies usually require an operating reserve.", False),
    ("SEC", "Growth: own Part 135 certificate and third-party management"),
    ("cert_delay", "Delay to own-certificate year set on Fleet Plan", "years", 1, 0, -1, NUM,
     "FAA certification typically 12-24 months after application; downside assumes a one-year slip.", False),
    ("cert_cost", "Certificate startup cost (year before)", "$", 450_000, 350_000, 275_000, USD,
     "Required staff hired early (Director of Ops, Chief Pilot, Director of Maintenance), manuals, consultants, proving runs.", False),
    ("cert_oh", "Own-certificate overhead (fixed)", "$/yr", 1_000_000, 850_000, 750_000, USD,
     "DO, DOM, ops control/dispatch, SMS, ARGUS/Wyvern audits, sales, accounting.", False),
    ("cert_oh_ac", "Own-certificate overhead per aircraft operated", "$/yr", 60_000, 50_000, 40_000, USD, "", False),
    ("tp_margin", "Net margin per third-party managed aircraft", "$/yr", 150_000, 200_000, 225_000, USD,
     "Management fee + charter commission minus direct staff cost. Estimate.", True),
    ("SEC", "Stop-loss rule (Henry, Oct 2026)"),
    ("stoploss", "Mature-year cash flow after debt that blocks buying more jets", "$/yr", -750_000, -750_000, -750_000, USD,
     "If one jet's mature-year cash flow after debt is worse than this, the model freezes the owned fleet at one jet. "
     "If the test fails, Jet #1 is also sold in the stop-loss year (see next line).", True),
    ("stoploss_year", "Stop-loss sale year if the test fails", "year", 2029, 2029, 2029, YEAR,
     "Henry: sell Jet #1 if trailing-12 cash after debt is worse than the threshold by mid-2029. Sold mid-year.", False),
    ("SEC", "Investor structure (GP/LP)"),
    ("gp_coinv", "GP (Henry + Chief Pilot) share of equity contributed", "%", 0.0, 0.0, 0.0, PCT,
     "Henry: no personal capital; investors fund the equity. Sweat equity is rewarded through the promote.", False),
    ("pref", "Investor preferred return (compounding)", "%", 0.08, 0.08, 0.08, PCT,
     "Common in real-asset private placements (6-10%).", True),
    ("lp_split", "Investor share of profits above the preferred return", "%", 0.80, 0.80, 0.80, PCT,
     "GP promote = the rest (20%). Common 70/30 to 80/20.", True),
    ("acq_fee", "GP acquisition fee (% of aircraft price)", "%", 0.01, 0.01, 0.01, PCT,
     "Paid by the company when each jet closes; covers sourcing and deal work. Common 1-2%.", False),
    ("SEC", "Exit and tax"),
    ("exit_year", "Exit year", "year", 2036, 2036, 2036, YEAR, "Henry: grow the fleet over 5-10 years, then exit. 2027-2036 horizon.", True),
    ("multiple", "EBITDA multiple on management platform", "x", 5.0, 6.0, 6.5, MULT,
     "Applied to third-party management margin only; owned fleet valued at market. Estimate.", True),
    ("sell_cost", "Fleet selling costs at exit", "%", 0.03, 0.02, 0.02, PCT, "Broker commissions, pre-sale inspections.", False),
    ("dep", "Annual decline in aircraft market value", "%", 0.08, 0.06, 0.045, PCT,
     "Estimate; Sovereign+ production ended, so values may soften faster than in-production types.", True),
    ("tax", "Investor ordinary tax rate (illustrative)", "%", 0.37, 0.37, 0.37, PCT,
     "Memo only. Investors can use losses only against passive income unless they materially participate. Consult a CPA.", False),
    ("cg_tax", "Investor capital gains rate incl. 3.8% NIIT (illustrative)", "%", 0.238, 0.238, 0.238, PCT,
     "Applied to platform (goodwill) value at exit.", False),
    ("bonus", "First-year bonus depreciation", "%", 1.0, 1.0, 1.0, PCT,
     "2025 law restored 100% bonus depreciation permanently; used aircraft qualify; >50% qualified business use. globalair.com", False),
]

AR = {}  # key -> row
row = 6
for item in ASSUMPTIONS:
    if item[0] == "SEC":
        section(A, row, item[1], 7)
        row += 1
        continue
    key, label, unit, d, b, u, fmt, note, lever = item
    AR[key] = row
    A.cell(row=row, column=1, value=label).font = BLACK
    A.cell(row=row, column=2, value=unit).font = SUB
    for col, v in zip("CDE", (d, b, u)):
        put(A, f"{col}{row}", v, fmt, fill=KEY_FILL if lever else None)
    put(A, f"F{row}", f"=CHOOSE($C$3,C{row},D{row},E{row})", fmt, bold=True)
    A.cell(row=row, column=7, value=note).font = SUB
    row += 1


def L(key):
    return f"Assumptions!$F${AR[key]}"


def a(key, col):
    return f"Assumptions!${col}${AR[key]}"


# ------------------------------------------------------------ Unit Economics
U = wb.create_sheet("Unit Economics")
setup(U, "Unit economics: one Sovereign+, mature year, on EJM's Part 135 certificate",
      "Today's dollars, no inflation. Shows whether a single financed jet carries itself.",
      {"A": 50, "B": 14, "C": 15, "D": 15, "E": 15, "F": 15})
header(U, 4, ["Line item", "", "Downside", "Base", "Upside", "Live"])
UR = {}


def urow(r, key, label, formula_fn, fmt=USD, bold=False, border=False):
    UR[key] = r
    U.cell(row=r, column=1, value=label).font = BOLD if bold else BLACK
    for col in "CDEF":
        c = put(U, f"{col}{r}", formula_fn(col), fmt, bold=bold)
        if border:
            c.border = TOP


r = 5
urow(r, "hours", "Billable hours", lambda c: f"={a('hours', c)}", NUM); r += 1
urow(r, "rate", "Blended rate per billable hour", lambda c: f"={a('mix', c)}*{a('rate_ovf', c)}+(1-{a('mix', c)})*{a('rate_ret', c)}"); r += 1
urow(r, "rev", "Charter revenue", lambda c: f"={c}{UR['hours']}*{c}{UR['rate']}", bold=True); r += 1
urow(r, "comm", "Less manager's commission", lambda c: f"=-{c}{UR['rev']}*{a('comm', c)}"); r += 1
urow(r, "net", "Net charter revenue to owner", lambda c: f"={c}{UR['rev']}+{c}{UR['comm']}", bold=True, border=True); r += 2
urow(r, "fh", "Flight hours incl. unpaid positioning", lambda c: f"={c}{UR['hours']}*(1+{a('ferry', c)})", NUM); r += 1
urow(r, "varhr", "Variable cost per flight hour",
     lambda c: f"={a('gph', c)}*{a('fuel', c)}*(1-{a('fuel_paid', c)})+({a('maint', c)}+{a('engine', c)}+{a('apu', c)})*(1+{a('maint_cont', c)})+{a('trip', c)}"); r += 1
urow(r, "var", "Variable costs", lambda c: f"=-{c}{UR['fh']}*{c}{UR['varhr']}"); r += 2
urow(r, "crew", "Crew salaries and benefits",
     lambda c: f"=-(({a('captains', c)}*{a('cap_sal', c)}+{a('fos', c)}*{a('fo_sal', c)})*(1+{a('benefits', c)})+{a('contract', c)})"); r += 1
urow(r, "train", "Recurrent training", lambda c: f"=-({a('captains', c)}+{a('fos', c)})*{a('train', c)}"); r += 1
urow(r, "hangar", "Hangar", lambda c: f"=-{a('hangar', c)}"); r += 1
urow(r, "ins", "Insurance", lambda c: f"=-{a('ins', c)}"); r += 1
urow(r, "mfee", "Management fee", lambda c: f"=-{a('mgmt_fee', c)}*12"); r += 1
urow(r, "other", "Other fixed", lambda c: f"=-{a('other', c)}"); r += 1
urow(r, "cont", "Contingency", lambda c: f"=-{a('conting', c)}"); r += 1
urow(r, "ga", "Company G&A", lambda c: f"=-{a('ga', c)}"); r += 1
urow(r, "fixed", "Total fixed costs", lambda c: f"=SUM({c}{UR['crew']}:{c}{UR['ga']})", bold=True, border=True); r += 2
urow(r, "ebitda", "EBITDA (operating cash flow before debt)",
     lambda c: f"={c}{UR['net']}+{c}{UR['var']}+{c}{UR['fixed']}", bold=True); r += 1
urow(r, "margin", "EBITDA margin", lambda c: f"=IF({c}{UR['rev']}=0,0,{c}{UR['ebitda']}/{c}{UR['rev']})", PCT); r += 2
urow(r, "loan", "Loan amount", lambda c: f"={a('price', c)}*{a('ltv', c)}"); r += 1
urow(r, "ds", "Annual debt service", lambda c: f"=-PMT({a('rate', c)},{a('amort', c)},-{c}{UR['loan']})"); r += 1
urow(r, "cfads", "Cash flow after debt service", lambda c: f"={c}{UR['ebitda']}+{c}{UR['ds']}", bold=True, border=True); r += 1
urow(r, "cfads_dry", "What-if: cash after debt if NetJets pays all fuel",
     lambda c: f"={c}{UR['cfads']}+{c}{UR['fh']}*{a('gph', c)}*{a('fuel', c)}*(1-{a('fuel_paid', c)})"); r += 1
urow(r, "int", "Of which: first-year interest", lambda c: f"=-{c}{UR['loan']}*{a('rate', c)}"); r += 1
urow(r, "decl", "First-year decline in aircraft value", lambda c: f"=-{a('price', c)}*{a('dep', c)}"); r += 1
urow(r, "econ", "Economic profit (EBITDA - interest - value decline)",
     lambda c: f"={c}{UR['ebitda']}+{c}{UR['int']}+{c}{UR['decl']}", bold=True); r += 2
urow(r, "contrib", "Contribution per billable hour",
     lambda c: f"={c}{UR['rate']}*(1-{a('comm', c)})-{c}{UR['varhr']}*(1+{a('ferry', c)})"); r += 1
urow(r, "be", "Breakeven billable hours (EBITDA = 0)",
     lambda c: f"=IF({c}{UR['contrib']}<=0,0,-{c}{UR['fixed']}/{c}{UR['contrib']})", NUM); r += 1
urow(r, "be_ds", "Breakeven billable hours incl. debt service",
     lambda c: f"=IF({c}{UR['contrib']}<=0,0,-({c}{UR['fixed']}+{c}{UR['ds']})/{c}{UR['contrib']})", NUM, bold=True); r += 1
urow(r, "req_rate", "NetJets rate needed for EBITDA = 0 at these hours",
     lambda c: f"=IF({a('mix', c)}=0,0,((-{c}{UR['var']}-{c}{UR['fixed']})/({c}{UR['hours']}*(1-{a('comm', c)}))-(1-{a('mix', c)})*{a('rate_ret', c)})/{a('mix', c)})"); r += 1
urow(r, "req_rate_ds", "NetJets rate needed to cover debt service too",
     lambda c: f"=IF({a('mix', c)}=0,0,((-{c}{UR['var']}-{c}{UR['fixed']}-{c}{UR['ds']})/({c}{UR['hours']}*(1-{a('comm', c)}))-(1-{a('mix', c)})*{a('rate_ret', c)})/{a('mix', c)})",
     bold=True); r += 1
urow(r, "reserve", "Cash reserve to hold", lambda c: f"=-{c}{UR['fixed']}/12*{a('wc_months', c)}"); r += 1
urow(r, "growth_ok", "Stop-loss test passed? (1 = may buy more jets)",
     lambda c: f"=IF({c}{UR['cfads']}>={a('stoploss', c)},1,0)", NUM, bold=True); r += 2

section(U, r, "Sensitivity: Live scenario, cash flow after debt service by billable hours", 6); r += 1
header(U, r, ["Billable hours", "", "EBITDA", "Debt service", "Cash after debt", "Per-hour margin"]); r += 1
for h in (300, 400, 500, 600, 700, 800, 900):
    put(U, f"A{r}", h, NUM)
    put(U, f"C{r}", f"=A{r}*$F${UR['contrib']}+$F${UR['fixed']}", USD)
    put(U, f"D{r}", f"=$F${UR['ds']}", USD)
    put(U, f"E{r}", f"=C{r}+D{r}", USD, bold=True)
    put(U, f"F{r}", f"=IF(A{r}=0,0,E{r}/A{r})", USD)
    r += 1

UE = q("Unit Economics")

# ---------------------------------------------------------------- Fleet Plan
F = wb.create_sheet("Fleet Plan")
setup(F, "Fleet plan by growth strategy", "Aircraft counts at year end (inputs). Owned = financed Sovereign+ jets. "
      "Third-party = other owners' jets managed on your certificate (only counts once you hold one). "
      "Owned growth stops at one jet if the stop-loss test fails.",
      {"A": 46, "B": 12, **{c: 10 for c in YCOLS}})
header(F, 4, ["Item", "Cert. year"] + [str(y) for y in YEARS])
FR = {}
r = 5
for code, name, _, desc in STRATEGIES:
    section(F, r, f"{name}: {desc}", 2 + len(YEARS)); r += 1
    F.cell(row=r, column=1, value="Owned aircraft, plan").font = BLACK
    put(F, f"B{r}", FLEET[code]["cert"], YEAR, fill=KEY_FILL)
    cert_r = r
    for i, c in enumerate(YCOLS):
        put(F, f"{c}{r}", FLEET[code]["owned"][i], NUM, fill=KEY_FILL)
    plan_r = r; r += 1
    F.cell(row=r, column=1, value="Third-party managed aircraft, plan").font = BLACK
    for i, c in enumerate(YCOLS):
        put(F, f"{c}{r}", FLEET[code]["tp"][i], NUM, fill=KEY_FILL)
    tp_r = r; r += 1
    F.cell(row=r, column=1, value="Owned aircraft bought (growth frozen at one jet if stop-loss fails)").font = BLACK
    for c in YCOLS:
        put(F, f"{c}{r}", f"=IF({UE}!$F${UR['growth_ok']}=1,{c}{plan_r},MIN({c}{plan_r},1))", NUM)
    bought_r = r; r += 1
    F.cell(row=r, column=1, value="Owned aircraft in fleet (after any stop-loss sale)").font = BOLD
    for i, c in enumerate(YCOLS):
        put(F, f"{c}{r}", f"=IF(AND({UE}!$F${UR['growth_ok']}=0,{2027 + i}>={L('stoploss_year')}),0,{c}{bought_r})", NUM, bold=True)
    FR[code] = {"owned": r, "bought": bought_r, "tp": tp_r, "cert": cert_r}
    r += 2


# --------------------------------------------------- Forecast + Debt per strategy
def build_strategy(code, name, suffix):
    dname, fname = f"Debt {suffix}", f"Forecast {suffix}"
    D = wb.create_sheet(dname)
    FC = wb.create_sheet(fname)
    fr = FR[code]
    owned_rng = f"'Fleet Plan'!$C${fr['bought']}:${LASTY}${fr['bought']}"
    sale_yr = f"IF({UE}!$F${UR['growth_ok']}=1,9999,{L('stoploss_year')})"

    # ---- debt schedule
    setup(D, f"Debt and aircraft value schedule: {name} strategy",
          "Each owned aircraft gets its own loan at the Live LTV and rate. Payments start in the purchase year (conservative).",
          {"A": 34, "B": 14, **{c: 13 for c in YCOLS}})
    header(D, 4, ["Item", ""] + [str(y) for y in YEARS])
    put(D, "A5", "Year", font=BOLD)
    for i, c in enumerate(YCOLS):
        put(D, f"{c}5", YEARS[i], YEAR, font=BOLD)
    DR = {k: [] for k in ("int", "prin", "end", "val", "payoff", "saleproc")}
    r = 7
    for k in range(1, MAX_SLOTS + 1):
        section(D, r, f"Aircraft #{k}", 2 + len(YEARS)); r += 1
        D.cell(row=r, column=1, value="Purchase year (2037 = not bought)").font = BLACK
        put(D, f"B{r}", f"=2027+COUNTIF({owned_rng},\"<{k}\")", YEAR)
        py = f"$B${r}"; r += 1
        D.cell(row=r, column=1, value="Price paid").font = BLACK
        put(D, f"B{r}", f"=IF({py}>2036,0,{L('price')}*(1+{L('price_chg')})^({py}-2027))", USD)
        pp = f"$B${r}"; r += 1
        D.cell(row=r, column=1, value="Loan").font = BLACK
        put(D, f"B{r}", f"={pp}*{L('ltv')}", USD)
        ln = f"$B${r}"; r += 1
        D.cell(row=r, column=1, value="Annual payment").font = BLACK
        put(D, f"B{r}", f"=IF({ln}=0,0,PMT({L('rate')},{L('amort')},-{ln}))", USD)
        pm = f"$B${r}"; r += 1
        rows = {}
        for nm, label in (("begin", "Beginning balance"), ("int", "Interest"), ("pay", "Payment"),
                          ("prin", "Principal"), ("payoff", "Loan payoff at stop-loss sale"),
                          ("end", "Ending balance"), ("val", "Market value (year end, while owned)"),
                          ("saleproc", "Stop-loss sale proceeds (net of selling costs)")):
            rows[nm] = r
            if nm in DR:
                DR[nm].append(r)
            D.cell(row=r, column=1, value=label).font = BLACK
            r += 1
        for i, c in enumerate(YCOLS):
            y = f"{c}$5"
            prev_end = "0" if i == 0 else f"{YCOLS[i-1]}{rows['end']}"
            put(D, f"{c}{rows['begin']}", f"=IF({y}={py},{ln},IF(AND({y}>{py},{y}<={sale_yr}),{prev_end},0))", USD)
            put(D, f"{c}{rows['int']}", f"={c}{rows['begin']}*{L('rate')}", USD)
            put(D, f"{c}{rows['pay']}", f"=IF({c}{rows['begin']}>0,MIN({pm},{c}{rows['begin']}+{c}{rows['int']}),0)", USD)
            put(D, f"{c}{rows['prin']}", f"={c}{rows['pay']}-{c}{rows['int']}", USD)
            put(D, f"{c}{rows['payoff']}", f"=IF({y}={sale_yr},{c}{rows['begin']}-{c}{rows['prin']},0)", USD)
            put(D, f"{c}{rows['end']}", f"={c}{rows['begin']}-{c}{rows['prin']}-{c}{rows['payoff']}", USD)
            put(D, f"{c}{rows['val']}", f"=IF(AND({y}>={py},{y}<{sale_yr}),{pp}*(1-{L('dep')})^({y}-{py}+1),0)", USD)
            put(D, f"{c}{rows['saleproc']}",
                f"=IF(AND({y}={sale_yr},{y}>={py}),{pp}*(1-{L('dep')})^({y}-{py}+0.5)*(1-{L('sell_cost')}),0)", USD)
        r += 1
    section(D, r, "Fleet totals", 2 + len(YEARS)); r += 1
    DT = {}
    for nm, label in (("int", "Total interest"), ("prin", "Total principal"), ("end", "Total debt outstanding"),
                      ("val", "Total fleet market value"), ("payoff", "Total loan payoff at stop-loss sale"),
                      ("saleproc", "Total stop-loss sale proceeds")):
        DT[nm] = r
        D.cell(row=r, column=1, value=label).font = BOLD
        for c in YCOLS:
            put(D, f"{c}{r}", "=" + "+".join(f"{c}{x}" for x in DR[nm]), USD, bold=True)
        r += 1
    DQ = q(dname)

    # ---- forecast
    setup(FC, f"10-year forecast: {name} strategy (Live scenario)",
          "Annual, nominal dollars. Owned jets fly on EJM's certificate until the own-certificate year.",
          {"A": 56, "B": 4, **{c: 13 for c in YCOLS}})
    put(FC, "A3", "=\"Scenario: \"&Assumptions!$D$3")
    header(FC, 4, ["Line item", ""] + [str(y) for y in YEARS])
    R = {}

    def frow(r, key, label, fn, fmt=USD, bold=False, border=False):
        R[key] = r
        FC.cell(row=r, column=1, value=label).font = BOLD if bold else BLACK
        for i, c in enumerate(YCOLS):
            prev = YCOLS[i - 1] if i else None
            cell = put(FC, f"{c}{r}", fn(c, prev, i), fmt, bold=bold)
            if border:
                cell.border = TOP

    r = 5
    frow(r, "year", "Year", lambda c, p, i: YEARS[i], YEAR, bold=True); r += 1
    frow(r, "idx", "Years from 2027", lambda c, p, i: f"={c}{R['year']}-2027", NUM); r += 1
    frow(r, "cert", "On EJM's certificate? (1 = yes; stays 1 if stop-loss ends an owned-only plan)",
         lambda c, p, i: f"=IF(OR({c}{R['year']}<'Fleet Plan'!$B${fr['cert']}+{L('cert_delay')},"
                         f"AND({UE}!$F${UR['growth_ok']}=0,SUM('Fleet Plan'!$C${fr['tp']}:${LASTY}${fr['tp']})=0)),1,0)", NUM); r += 1
    frow(r, "cidx", "Cost inflation index", lambda c, p, i: f"=(1+{L('infl')})^{c}{R['idx']}", '0.000'); r += 2

    section(FC, r, "Fleet and activity", 12); r += 1
    frow(r, "owned", "Owned aircraft (year end)", lambda c, p, i: f"='Fleet Plan'!{c}{fr['owned']}", NUM); r += 1
    frow(r, "acy", "Owned aircraft-years in service",
         lambda c, p, i: f"={c}{R['owned']}*{L('months_y1')}/12" if i == 0 else f"=({p}{R['owned']}+{c}{R['owned']})/2", NUM1); r += 1
    frow(r, "util", "Utilization factor", lambda c, p, i: f"=IF({c}{R['idx']}=0,{L('ramp1')},IF({c}{R['idx']}=1,{L('ramp2')},1))", PCT); r += 1
    frow(r, "hrs", "Billable hours", lambda c, p, i: f"={c}{R['acy']}*{L('hours')}*{c}{R['util']}", NUM); r += 1
    frow(r, "fh", "Flight hours incl. positioning", lambda c, p, i: f"={c}{R['hrs']}*(1+{L('ferry')})", NUM); r += 1
    frow(r, "tp", "Third-party managed aircraft (year end)",
         lambda c, p, i: f"='Fleet Plan'!{c}{fr['tp']}*(1-{c}{R['cert']})", NUM); r += 1
    frow(r, "tpavg", "Third-party managed aircraft (average)",
         lambda c, p, i: "=0" if i == 0 else f"=({p}{R['tp']}+{c}{R['tp']})/2", NUM1); r += 2

    section(FC, r, "Revenue", 12); r += 1
    frow(r, "r_ovf", "Overflow / subcontract revenue",
         lambda c, p, i: f"={c}{R['hrs']}*{L('mix')}*{L('rate_ovf')}*(1+{L('rate_g')})^{c}{R['idx']}"); r += 1
    frow(r, "r_ret", "Broker / retail charter revenue",
         lambda c, p, i: f"={c}{R['hrs']}*(1-{L('mix')})*{L('rate_ret')}*(1+{L('rate_g')})^{c}{R['idx']}"); r += 1
    frow(r, "r_tp", "Third-party management margin (own certificate)",
         lambda c, p, i: f"={c}{R['tpavg']}*{L('tp_margin')}*{c}{R['cidx']}"); r += 1
    frow(r, "rev", "Total revenue", lambda c, p, i: f"=SUM({c}{R['r_ovf']}:{c}{R['r_tp']})", bold=True, border=True); r += 2

    section(FC, r, "Operating costs", 12); r += 1
    frow(r, "c_var", "Variable flight costs (fuel, maintenance + contingency, engines, APU, trip)",
         lambda c, p, i: f"={c}{R['fh']}*({L('gph')}*{L('fuel')}*(1-{L('fuel_paid')})+({L('maint')}+{L('engine')}+{L('apu')})*(1+{L('maint_cont')})+{L('trip')})*{c}{R['cidx']}"); r += 1
    frow(r, "c_crew", "Crew salaries, benefits, contract pilots",
         lambda c, p, i: f"={c}{R['acy']}*(({L('captains')}*{L('cap_sal')}+{L('fos')}*{L('fo_sal')})*(1+{L('benefits')})+{L('contract')})*{c}{R['cidx']}"); r += 1
    frow(r, "c_train", "Recurrent training",
         lambda c, p, i: f"={c}{R['acy']}*({L('captains')}+{L('fos')})*{L('train')}*{c}{R['cidx']}"); r += 1
    frow(r, "c_fix", "Hangar, insurance, other fixed, contingency",
         lambda c, p, i: f"={c}{R['acy']}*({L('hangar')}+{L('ins')}+{L('other')}+{L('conting')})*{c}{R['cidx']}"); r += 1
    frow(r, "c_mfee", "Management fees (on manager's certificate)",
         lambda c, p, i: f"={c}{R['cert']}*{c}{R['acy']}*{L('mgmt_fee')}*12*{c}{R['cidx']}"); r += 1
    frow(r, "c_comm", "Manager's charter commission",
         lambda c, p, i: f"={c}{R['cert']}*({c}{R['r_ovf']}+{c}{R['r_ret']})*{L('comm')}"); r += 1
    frow(r, "c_ga", "Company G&A (before own certificate)",
         lambda c, p, i: f"={c}{R['cert']}*IF({c}{R['acy']}>0,1,0)*{L('ga')}*{c}{R['cidx']}"); r += 1
    frow(r, "c_cert", "Own-certificate overhead",
         lambda c, p, i: f"=(1-{c}{R['cert']})*({L('cert_oh')}+{L('cert_oh_ac')}*({c}{R['acy']}+{c}{R['tpavg']}))*{c}{R['cidx']}"); r += 1
    frow(r, "opex", "Total operating costs", lambda c, p, i: f"=SUM({c}{R['c_var']}:{c}{R['c_cert']})", bold=True, border=True); r += 2

    frow(r, "ebitda", "EBITDA", lambda c, p, i: f"={c}{R['rev']}-{c}{R['opex']}", bold=True); r += 1
    frow(r, "margin", "EBITDA margin", lambda c, p, i: f"=IF({c}{R['rev']}=0,0,{c}{R['ebitda']}/{c}{R['rev']})", PCT); r += 1
    frow(r, "onetime", "One-time: own-certificate startup",
         lambda c, p, i: f"=IF(AND({c}{R['year']}='Fleet Plan'!$B${fr['cert']}+{L('cert_delay')}-1,"
                         f"OR({UE}!$F${UR['growth_ok']}=1,SUM('Fleet Plan'!$C${fr['tp']}:${LASTY}${fr['tp']})>0)),{L('cert_cost')},0)"); r += 1
    frow(r, "interest", "Interest", lambda c, p, i: f"={DQ}!{c}{DT['int']}"); r += 1
    frow(r, "principal", "Principal repayment", lambda c, p, i: f"={DQ}!{c}{DT['prin']}"); r += 1
    frow(r, "saleproc", "Stop-loss aircraft sale proceeds", lambda c, p, i: f"={DQ}!{c}{DT['saleproc']}"); r += 1
    frow(r, "payoff", "Loan payoff at stop-loss sale", lambda c, p, i: f"={DQ}!{c}{DT['payoff']}"); r += 1
    frow(r, "cfo", "Cash flow after debt service and any stop-loss sale",
         lambda c, p, i: f"={c}{R['ebitda']}-{c}{R['onetime']}-{c}{R['interest']}-{c}{R['principal']}+{c}{R['saleproc']}-{c}{R['payoff']}",
         bold=True, border=True); r += 2

    section(FC, r, "Aircraft purchases and investor equity", 12); r += 1
    frow(r, "newac", "Aircraft purchased", lambda c, p, i: f"={c}{R['owned']}" if i == 0 else f"={c}{R['owned']}-{p}{R['owned']}", NUM); r += 1
    frow(r, "price", "Price per aircraft", lambda c, p, i: f"={L('price')}*(1+{L('price_chg')})^{c}{R['idx']}"); r += 1
    frow(r, "capex", "Aircraft purchases (total price)", lambda c, p, i: f"={c}{R['newac']}*{c}{R['price']}"); r += 1
    frow(r, "newdebt", "Less: new loans", lambda c, p, i: f"=-{c}{R['capex']}*{L('ltv')}"); r += 1
    frow(r, "acqc", "Acquisition, onboarding, refurb, training, squawk reserve",
         lambda c, p, i: f"={c}{R['newac']}*({L('prebuy')}+{L('legal')}+{L('onboard')}+{L('refurb')}+{L('init_train')}+{L('squawk')})*{c}{R['cidx']}"); r += 1
    frow(r, "acqfee", "GP acquisition fee", lambda c, p, i: f"={c}{R['capex']}*{L('acq_fee')}"); r += 1
    frow(r, "setup", "Company setup and investor offering costs",
         lambda c, p, i: f"={L('entity')}+{L('offering')}" if i == 0 else "=0"); r += 1
    frow(r, "acqeq", "Equity needed for purchases", lambda c, p, i: f"=SUM({c}{R['capex']}:{c}{R['setup']})", bold=True, border=True); r += 2
    frow(r, "bcash", "Beginning cash", lambda c, p, i: "=0"); r += 1
    frow(r, "mincash", "Minimum cash reserve",
         lambda c, p, i: f"={c}{R['owned']}*{UE}!$F${UR['reserve']}*{c}{R['cidx']}"); r += 1
    frow(r, "neweq", "Equity contributed (investors + GP)",
         lambda c, p, i: f"=MAX(0,{c}{R['mincash']}-({c}{R['bcash']}+{c}{R['cfo']}-{c}{R['acqeq']}))", bold=True); r += 1
    frow(r, "ecash", "Ending cash", lambda c, p, i: f"={c}{R['bcash']}+{c}{R['cfo']}-{c}{R['acqeq']}+{c}{R['neweq']}", bold=True, border=True); r += 1
    for i, c in enumerate(YCOLS):  # beginning cash = prior ending cash
        if i:
            put(FC, f"{c}{R['bcash']}", f"={YCOLS[i-1]}{R['ecash']}", USD)
    frow(r, "cumeq", "Cumulative equity contributed",
         lambda c, p, i: f"={c}{R['neweq']}" if i == 0 else f"={p}{R['cumeq']}+{c}{R['neweq']}", bold=True); r += 2

    section(FC, r, "Balance sheet view (year end)", 12); r += 1
    frow(r, "fleetval", "Fleet market value", lambda c, p, i: f"={DQ}!{c}{DT['val']}"); r += 1
    frow(r, "debt", "Debt outstanding", lambda c, p, i: f"={DQ}!{c}{DT['end']}"); r += 1
    frow(r, "nav", "Net asset value (fleet + cash - debt)", lambda c, p, i: f"={c}{R['fleetval']}+{c}{R['ecash']}-{c}{R['debt']}", bold=True, border=True); r += 1
    frow(r, "ltvnow", "Loan-to-value on fleet", lambda c, p, i: f"=IF({c}{R['fleetval']}=0,0,{c}{R['debt']}/{c}{R['fleetval']})", PCT); r += 2

    section(FC, r, "Investor tax memo (passes through on K-1s; consult a CPA)", 12); r += 1
    frow(r, "taxdep", "Tax depreciation (100% bonus in purchase year)", lambda c, p, i: f"={c}{R['capex']}*{L('bonus')}"); r += 1
    frow(r, "taxinc", "Taxable income (loss) allocated to investors",
         lambda c, p, i: f"={c}{R['ebitda']}-{c}{R['onetime']}-{c}{R['interest']}-{c}{R['acqfee']}-{c}{R['setup']}-{c}{R['taxdep']}+{c}{R['saleproc']}"); r += 1
    frow(r, "taxfx", "Tax saved (owed) if investors can use losses as they arise",
         lambda c, p, i: f"=-{c}{R['taxinc']}*{L('tax')}", bold=True); r += 1
    FC.cell(row=r, column=1, value="Losses are passive for most investors: usable only against other passive income, "
            "otherwise carried forward and released at sale. Depreciation is recaptured as ordinary income when jets are sold.").font = SUB
    r += 2

    # ---- exit and waterfall (single values in column C)
    section(FC, r, "Exit and investor waterfall", 12); r += 1
    yr_rng = f"$C${R['year']}:${LASTY}${R['year']}"

    def at_exit(key):
        return f"INDEX($C${R[key]}:${LASTY}${R[key]},MATCH({L('exit_year')},{yr_rng},0))"

    E = {}
    exit_items = [
        ("year", "Exit year", f"={L('exit_year')}", YEAR),
        ("fleet", "Fleet market value", f"={at_exit('fleetval')}", USD),
        ("sellc", "Less selling costs", "=-C{fleet}*" + L('sell_cost'), USD),
        ("tpm", "Third-party management margin in exit year", f"={at_exit('r_tp')}", USD),
        ("plat", "Platform value (margin x multiple)", "=C{tpm}*" + L('multiple'), USD),
        ("ev", "Enterprise value", "=C{fleet}+C{sellc}+C{plat}", USD),
        ("debt", "Less debt outstanding", f"=-{at_exit('debt')}", USD),
        ("cash", "Plus cash on hand", f"={at_exit('ecash')}", USD),
        ("P", "Net exit proceeds (before investor taxes)", "=C{ev}+C{debt}+C{cash}", USD),
        ("C", "Total equity contributed through exit",
         f"=SUMPRODUCT(({yr_rng}<={L('exit_year')})*$C${R['neweq']}:${LASTY}${R['neweq']})", USD),
        ("prefamt", "Accrued preferred return",
         f"=SUMPRODUCT(({yr_rng}<={L('exit_year')})*$C${R['neweq']}:${LASTY}${R['neweq']}*((1+{L('pref')})^({L('exit_year')}-{yr_rng})-1))", USD),
        ("t1", "Tier 1: return of capital", "=MIN(MAX(C{P},0),C{C})", USD),
        ("t2", "Tier 2: preferred return", "=MIN(MAX(C{P},0)-C{t1},C{prefamt})", USD),
        ("t3", "Tier 3: remaining profit to split", "=MAX(C{P},0)-C{t1}-C{t2}", USD),
        ("lp", "Investors (LPs) receive",
         f"=(C{{t1}}+C{{t2}}+C{{t3}}*{L('lp_split')})*(1-{L('gp_coinv')})", USD),
        ("gp", "GP (Henry + Chief Pilot) receives: promote + co-invest share",
         f"=(C{{t1}}+C{{t2}}+C{{t3}}*{L('lp_split')})*{L('gp_coinv')}+C{{t3}}*(1-{L('lp_split')})", USD),
        ("gpfees", "GP acquisition fees earned through exit",
         f"=SUMPRODUCT(({yr_rng}<={L('exit_year')})*$C${R['acqfee']}:${LASTY}${R['acqfee']})", USD),
        ("lpmoic", "Investor multiple on capital (MOIC)", f"=IF(C{{C}}=0,0,C{{lp}}/(C{{C}}*(1-{L('gp_coinv')})))", MULT),
        ("lpirr", "Investor IRR (pre-tax)", "", PCT),
        ("lpirr_tax", "Investor IRR after tax, if losses usable as they arise", "", PCT),
        ("lpmoic_tax", "Investor after-tax multiple, if losses usable (tax savings counted as cash)",
         f"=IF(C{{C}}=0,0,(C{{lp}}+SUMPRODUCT(({yr_rng}<={L('exit_year')})*$C${R['taxfx']}:${LASTY}${R['taxfx']})*(1-{L('gp_coinv')})+C{{exittax}})/(C{{C}}*(1-{L('gp_coinv')})))",
         MULT),
        ("projirr", "Project IRR, all equity before promote (pre-tax)", "", PCT),
        ("peak", "Peak cumulative equity raised", f"=MAX($C${R['cumeq']}:${LASTY}${R['cumeq']})", USD),
        ("exittax", "Investor tax at exit (recapture + capital gains), illustrative",
         f"=-(MAX(0,C{{fleet}}+C{{sellc}})*{L('tax')}+C{{plat}}*{L('cg_tax')})*(1-{L('gp_coinv')})", USD),
    ]
    for key, label, _, _ in exit_items:
        E[key] = r
        r += 1
    for key, label, f, fmt in exit_items:
        rr = E[key]
        FC.cell(row=rr, column=1, value=label).font = BOLD if key in ("P", "lp", "gp", "lpmoic", "lpirr", "lpirr_tax") else BLACK
        if f:
            put(FC, f"C{rr}", f.format(**{k: v for k, v in E.items()}), fmt, bold=key in ("P", "lp", "gp", "lpmoic"))
    r += 1

    section(FC, r, "Cash flows for IRR", 12); r += 1
    frow(r, "exitflag", "Exit year flag", lambda c, p, i: f"=IF({c}{R['year']}={L('exit_year')},1,0)", NUM); r += 1
    frow(r, "f_proj", "All equity: contributions (-) and exit proceeds (+)",
         lambda c, p, i: f"=IF({c}{R['year']}>{L('exit_year')},0,-{c}{R['neweq']}+{c}{R['exitflag']}*$C${E['P']})"); r += 1
    frow(r, "f_lp", "Investors (LPs), pre-tax",
         lambda c, p, i: f"=IF({c}{R['year']}>{L('exit_year')},0,-{c}{R['neweq']}*(1-{L('gp_coinv')})+{c}{R['exitflag']}*$C${E['lp']})",
         bold=True); r += 1
    frow(r, "f_lptax", "Investors (LPs), after tax if losses usable",
         lambda c, p, i: f"=IF({c}{R['year']}>{L('exit_year')},0,{c}{R['f_lp']}+{c}{R['taxfx']}*(1-{L('gp_coinv')})+{c}{R['exitflag']}*$C${E['exittax']})"); r += 1
    for key, rowkey in (("lpirr", "f_lp"), ("lpirr_tax", "f_lptax"), ("projirr", "f_proj")):
        put(FC, f"C{E[key]}", f"=IFERROR(IRR($C${R[rowkey]}:${LASTY}${R[rowkey]},0.1),IFERROR(IRR($C${R[rowkey]}:${LASTY}${R[rowkey]},-0.1),\"n/a\"))", PCT, bold=key != "projirr")

    FC.freeze_panes = "C5"
    D.freeze_panes = "C5"
    return fname, R, E


SR = {}
for code, name, suffix, _ in STRATEGIES:
    SR[code] = build_strategy(code, name, suffix)

# ------------------------------------------------------- Financing Scenarios
FS = wb.create_sheet("Financing Scenarios", 1)
setup(FS, "Financing scenarios: first aircraft, Live scenario",
      "How leverage changes cash needed and annual cash flow. Change the LTV row (yellow) to test other structures.",
      {"A": 50, "B": 4, **{get_column_letter(3 + i): 14 for i in range(6)}})
LTVS = [0.60, 0.70, 0.75, 0.80, 0.85, 0.90]
LC = [get_column_letter(3 + i) for i in range(len(LTVS))]
header(FS, 4, ["", ""] + ["Option " + str(i + 1) for i in range(len(LTVS))])
FSR = {}


def fsrow(r, key, label, fn, fmt=USD, bold=False, border=False):
    FSR[key] = r
    FS.cell(row=r, column=1, value=label).font = BOLD if bold else BLACK
    for col in LC:
        c = put(FS, f"{col}{r}", fn(col), fmt, bold=bold)
        if border:
            c.border = TOP


r = 5
FS.cell(row=r, column=1, value="Loan-to-value").font = BOLD
for col, v in zip(LC, LTVS):
    put(FS, f"{col}{r}", v, PCT, fill=KEY_FILL)
FSR["ltv"] = r; r += 1
fsrow(r, "price", "Purchase price", lambda c: f"={L('price')}"); r += 1
fsrow(r, "loan", "Loan", lambda c: f"={c}{FSR['price']}*{c}{FSR['ltv']}"); r += 1
fsrow(r, "down", "Down payment", lambda c: f"={c}{FSR['price']}-{c}{FSR['loan']}"); r += 1
fsrow(r, "acq", "Acquisition, onboarding, setup, offering costs, GP fee",
      lambda c: f"={L('prebuy')}+{L('legal')}+{L('onboard')}+{L('refurb')}+{L('init_train')}+{L('squawk')}+{L('entity')}+{L('offering')}+{L('price')}*{L('acq_fee')}"); r += 1
fsrow(r, "res", "Cash reserve", lambda c: f"={UE}!$F${UR['reserve']}"); r += 1
fsrow(r, "cash", "Total equity to raise for launch", lambda c: f"=SUM({c}{FSR['down']}:{c}{FSR['res']})", bold=True, border=True); r += 2
fsrow(r, "ds", "Annual debt service", lambda c: f"=PMT({L('rate')},{L('amort')},-{c}{FSR['loan']})"); r += 1
fsrow(r, "dsm", "Monthly loan payment", lambda c: f"={c}{FSR['ds']}/12"); r += 1
fsrow(r, "ebitda", "Mature-year EBITDA (Unit Economics, Live)", lambda c: f"={UE}!$F${UR['ebitda']}"); r += 1
fsrow(r, "cf", "Cash flow after debt service", lambda c: f"={c}{FSR['ebitda']}-{c}{FSR['ds']}", bold=True, border=True); r += 1
fsrow(r, "dscr", "Debt service coverage (EBITDA / debt service)",
      lambda c: f"=IF({c}{FSR['ds']}=0,0,{c}{FSR['ebitda']}/{c}{FSR['ds']})", MULT); r += 1
fsrow(r, "coc", "Cash-on-cash return", lambda c: f"=IF({c}{FSR['cash']}=0,0,{c}{FSR['cf']}/{c}{FSR['cash']})", PCT); r += 1
fsrow(r, "be", "Breakeven billable hours incl. debt service",
      lambda c: f"=IF({UE}!$F${UR['contrib']}<=0,0,(-{UE}!$F${UR['fixed']}+{c}{FSR['ds']})/{UE}!$F${UR['contrib']})",
      NUM, bold=True); r += 1
fsrow(r, "taxloss", "First-year tax loss from 100% bonus depreciation (approx.)",
      lambda c: f"={L('price')}*{L('bonus')}"); r += 1
fsrow(r, "lossratio", "Year-1 depreciation per $1 of investor equity",
      lambda c: f"=IF({c}{FSR['cash']}=0,0,{c}{FSR['taxloss']}/{c}{FSR['cash']})", MULT, bold=True); r += 2
FS.cell(row=r, column=1, value="Lender reality check").font = BOLD; r += 1
for note in (
    "Credit-based jet loans typically need 15-20% down (80-85% LTV); asset-based loans 30-40% down (AOPA via flycraft.com).",
    "An investor-owned SPV with no wealthy guarantor will usually be underwritten as asset-based: expect 60-70% LTV unless an investor guarantees.",
    "85-90% leverage usually means a personal guarantee, strong outside liquidity, or a second-lien / seller note on top of the senior loan.",
    "Higher leverage raises depreciation per equity dollar (attractive to investors) but deepens annual cash losses.",
):
    FS.cell(row=r, column=1, value=note).font = SUB
    r += 1
r += 1
section(FS, r, "Cash flow after debt service: billable hours (rows) by LTV (columns)", 8); r += 1
header(FS, r, ["Billable hours", ""] + [None] * len(LC))
for col in LC:
    put(FS, f"{col}{r}", f"={col}{FSR['ltv']}", PCT, font=HDR_FONT, fill=HDR_FILL)
r += 1
for h in (400, 500, 600, 700, 800, 900):
    put(FS, f"A{r}", h, NUM)
    for col in LC:
        put(FS, f"{col}{r}",
            f"=$A{r}*{UE}!$F${UR['contrib']}+{UE}!$F${UR['fixed']}-{col}${FSR['ds']}", USD)
    r += 1

# --------------------------------------------------------------- Startup Budget
S = wb.create_sheet("Startup Budget", 1)
setup(S, "Startup budget: first aircraft (Live scenario)", "Uses and sources of cash to get Jet #1 flying revenue trips.",
      {"A": 52, "B": 4, "C": 16, "D": 60})
header(S, 4, ["Uses of cash", "", "Amount", "Note"])
SB = [
    ("Purchase price", f"={L('price')}", "Assumptions"),
    ("Pre-buy inspection", f"={L('prebuy')}", ""),
    ("Closing legal, escrow, title, registration", f"={L('legal')}", ""),
    ("Part 135 onboarding, conformity, Wi-Fi/avionics", f"={L('onboard')}", ""),
    ("Interior/paint refresh", f"={L('refurb')}", ""),
    ("Initial crew training and hiring", f"={L('init_train')}", ""),
    ("Post-purchase squawk and maintenance reserve", f"={L('squawk')}", "One-time, conservative; unspent cash stays in the company"),
    ("Company formation, legal, accounting, insurance setup", f"={L('entity')}", ""),
    ("Investor offering costs (securities counsel, PPM, Form D)", f"={L('offering')}", ""),
    ("GP acquisition fee", f"={L('price')}*{L('acq_fee')}", "Paid to Henry's GP entity at closing"),
    ("Cash reserve (operating)", f"={UE}!$F${UR['reserve']}", "Months of fixed cost per Assumptions"),
]
r = 5
for label, f, note in SB:
    S.cell(row=r, column=1, value=label).font = BLACK
    put(S, f"C{r}", f, USD)
    S.cell(row=r, column=4, value=note).font = SUB
    r += 1
S.cell(row=r, column=1, value="Total uses").font = BOLD
put(S, f"C{r}", f"=SUM(C5:C{r-1})", USD, bold=True).border = TOP
tot = r
r += 2
header(S, r, ["Sources of cash", "", "Amount", "Note"]); r += 1
S.cell(row=r, column=1, value="Aircraft loan").font = BLACK
put(S, f"C{r}", f"={L('price')}*{L('ltv')}", USD)
S.cell(row=r, column=4, value="Senior secured loan at Live LTV").font = SUB
loan_r = r; r += 1
S.cell(row=r, column=1, value="Investor equity (LP units)").font = BLACK
put(S, f"C{r}", f"=(C{tot}-C{loan_r})*(1-{L('gp_coinv')})", USD)
lp_r = r; r += 1
S.cell(row=r, column=1, value="GP co-investment").font = BLACK
put(S, f"C{r}", f"=(C{tot}-C{loan_r})*{L('gp_coinv')}", USD)
gp_r = r; r += 1
S.cell(row=r, column=1, value="Total sources").font = BOLD
put(S, f"C{r}", f"=C{loan_r}+C{lp_r}+C{gp_r}", USD, bold=True).border = TOP
r += 2
S.cell(row=r, column=1, value="Equity as % of total uses").font = BLACK
put(S, f"C{r}", f"=(C{lp_r}+C{gp_r})/C{tot}", PCT)
r += 2
fq = q(SR["hyb"][0])
Rh = SR["hyb"][1]
col28 = YCOLS[1]
header(S, r, ["First full year operating budget (2028)", "", "Amount", "Note"]); r += 1
for label, key in (("Charter revenue", None), ("Variable flight costs", "c_var"), ("Crew", "c_crew"),
                   ("Training", "c_train"), ("Hangar, insurance, other fixed", "c_fix"), ("Management fees", "c_mfee"),
                   ("Manager's commission", "c_comm"), ("Company G&A", "c_ga"), ("EBITDA", "ebitda"),
                   ("Debt service", None), ("Cash flow after debt service", "cfo")):
    S.cell(row=r, column=1, value=label).font = BOLD if label in ("EBITDA", "Cash flow after debt service") else BLACK
    if label == "Charter revenue":
        f = f"={fq}!{col28}{Rh['r_ovf']}+{fq}!{col28}{Rh['r_ret']}"
    elif label == "Debt service":
        f = f"=-({fq}!{col28}{Rh['interest']}+{fq}!{col28}{Rh['principal']})"
    elif key in ("ebitda", "cfo"):
        f = f"={fq}!{col28}{Rh[key]}"
    else:
        f = f"=-{fq}!{col28}{Rh[key]}"
    put(S, f"C{r}", f, USD, bold=label in ("EBITDA", "Cash flow after debt service"))
    S.cell(row=r, column=4, value="One jet; identical in all three strategies in 2028" if label == "Charter revenue" else "").font = SUB
    r += 1

# ------------------------------------------------------------------ Dashboard
Z = wb.create_sheet("Dashboard", 0)
setup(Z, "Jet overflow charter business: dashboard", "Change the scenario on the Assumptions sheet (cell C3). All figures update.",
      {"A": 50, "B": 4, "C": 17, "D": 17, "E": 17, "F": 44})
put(Z, "A3", "=\"Scenario shown: \"&Assumptions!$D$3")
header(Z, 5, ["One jet, mature year", "", "Value", "", "", "What it means"])
DASH = [
    ("Purchase price, Jet #1", f"={L('price')}", USD, ""),
    ("Equity to raise for launch", f"='Startup Budget'!C{lp_r}+'Startup Budget'!C{gp_r}", USD, "Down payment + startup costs + reserve"),
    ("Aircraft loan", f"='Startup Budget'!C{loan_r}", USD, ""),
    ("EBITDA, one jet", f"={UE}!F{UR['ebitda']}", USD, "Before debt"),
    ("Annual debt service, one jet", f"=-{UE}!F{UR['ds']}", USD, ""),
    ("Cash flow after debt, one jet", f"={UE}!F{UR['cfads']}", USD, "Negative = investors fund the gap"),
    ("What-if: cash after debt if NetJets pays fuel", f"={UE}!F{UR['cfads_dry']}", USD, "Ask EJM whether the rate is wet or dry"),
    ("Breakeven billable hours (EBITDA)", f"={UE}!F{UR['be']}", NUM, "0 = each extra hour loses money"),
    ("Breakeven billable hours incl. debt", f"={UE}!F{UR['be_ds']}", NUM, "Over ~1,000 hrs means hours alone can't fix it"),
    ("NetJets rate needed, EBITDA breakeven", f"={UE}!F{UR['req_rate']}", USD, "Per billable hour, at the assumed hours"),
    ("NetJets rate needed, incl. debt service", f"={UE}!F{UR['req_rate_ds']}", USD, "Key number to get from EJM"),
    ("Stop-loss test passed? (1 = yes)", f"={UE}!F{UR['growth_ok']}", NUM, "0 = owned fleet frozen at one jet"),
]
r = 6
for label, f, fmt, note in DASH:
    Z.cell(row=r, column=1, value=label).font = BLACK
    put(Z, f"C{r}", f, fmt, bold=True)
    Z.cell(row=r, column=6, value=note).font = SUB
    r += 1
r += 1
header(Z, r, ["Growth strategy comparison (2027-2036)", ""] + [s[1] for s in STRATEGIES] + ["Note"]); r += 1
COMP = [
    ("Owned aircraft at exit", "owned", None, NUM, ""),
    ("Third-party managed aircraft at exit", "tp", None, NUM, ""),
    ("Revenue in exit year", "rev", None, USD, ""),
    ("EBITDA in exit year", "ebitda", None, USD, ""),
    ("Debt outstanding at exit", "debt", None, USD, ""),
    ("Peak equity to raise", None, "peak", USD, "Total investor capital needed"),
    ("Net exit proceeds", None, "P", USD, "Before investor taxes"),
    ("Investors receive (LP)", None, "lp", USD, "After 8% pref and 80/20 split"),
    ("Investor multiple (MOIC)", None, "lpmoic", MULT, ""),
    ("Investor IRR, pre-tax", None, "lpirr", PCT, ""),
    ("Investor after-tax multiple, if losses usable", None, "lpmoic_tax", MULT, "Counts depreciation tax savings, net of recapture"),
    ("Investor IRR, after tax if losses usable", None, "lpirr_tax", PCT, "n/a = year-1 tax savings exceed the cash invested"),
    ("GP promote (Henry + Chief Pilot)", None, "gp", USD, "Sweat-equity payoff at exit"),
    ("GP acquisition fees", None, "gpfees", USD, ""),
]
for label, rkey, ekey, fmt, note in COMP:
    Z.cell(row=r, column=1, value=label).font = BLACK
    for j, (code, *_rest) in enumerate(STRATEGIES):
        fname, R, E = SR[code]
        col = get_column_letter(3 + j)
        if rkey:
            f = (f"=INDEX({q(fname)}!$C${R[rkey]}:${LASTY}${R[rkey]},"
                 f"MATCH({L('exit_year')},{q(fname)}!$C${R['year']}:${LASTY}${R['year']},0))")
        else:
            f = f"={q(fname)}!$C${E[ekey]}"
        put(Z, f"{col}{r}", f, fmt, bold=True)
    Z.cell(row=r, column=6, value=note).font = SUB
    r += 1
r += 1
Z.cell(row=r, column=1, value="Legend").font = BOLD; r += 1
for txt, font in (("Blue text: input you can change", BLUE), ("Black text: formula", BLACK),
                  ("Green text: link to another sheet", GREEN)):
    Z.cell(row=r, column=1, value=txt).font = font
    r += 1
Z.cell(row=r, column=1, value="Yellow fill: key lever").fill = KEY_FILL
r += 1
Z.cell(row=r, column=1, value="All estimates must be verified with quotes from lenders, EJM, insurers and an aircraft broker.").font = SUB

wb.move_sheet("Assumptions", offset=-(len(wb.sheetnames) - 2))
wb.move_sheet("Fleet Plan", offset=-(wb.sheetnames.index("Fleet Plan") - 2))
for ws in wb.worksheets:
    if ws.title == "Fleet Plan":
        ws.freeze_panes = "C5"

wb.save(OUT)
print("saved", OUT, wb.sheetnames)
